import docx
from docx2pdf import convert
import openpyxl
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication


# 生成对应的邀请函，并转存pdf格式
def get_invitation(name):
    doc = docx.Document("template.docx")
    for para in doc.paragraphs:
        if '<name>' in para.text:
            for run in para.runs:
                if '<name>' in run.text:
                    run.text = run.text.replace('<name>', name)
        doc.save(f'./邀请函/{name}.docx')
    convert(f"./邀请函/{name}.docx")


smtp = smtplib.SMTP(host="smtp.qq.com", port=587)
smtp.login('235977@qq.com', "ruybefkipoo")


def send_email(name, email):
    msg = MIMEMultipart()
    msg["subject"] = f"您好，{name}，您的邀请函！"
    msg["from"] = "2352180977@qq.com"
    msg["to"] = email

    html_content = f"""
    <html>
        <body>
                <p>您好：{name}<br>
                    <b>欢迎加入Python进阶者学习交流群，请在附件中查收您的门票~</b><br>
                    点击这里了解更多：<a href="https://www.pdcfighting.com">演唱会主页</a>
                </p>
        </body>
    </html>
    """
    html_part = MIMEText(html_content, "html")
    msg.attach(html_part)
    with open(f"./邀请函/{name}.pdf", "rb") as f:
        doc_part = MIMEApplication(f.read())
        doc_part.add_header("Content-Disposition", "attachment", filename=name)
        # 把附件添加到邮件中
        msg.attach(doc_part)
        # 发送前面准备好的邮件
        smtp.send_message(msg)
        # 如果放到外边登录，这里就不用退出服务器连接，所以注释掉了
        # smtp.quit()


def get_username_email():
    workbook = openpyxl.load_workbook("names.xlsx")
    worksheet = workbook.active
    for index, row in enumerate(worksheet.rows):
        if index > 0:
            name = row[0].value
            email = row[3].value
            # print(name, email)
            # print(f"{name}邀请函正在生成...")
            # get_invitation(name)
            send_email(name, email)


if __name__ == '__main__':
    get_username_email()
    # get_invitation('Python进阶者')
